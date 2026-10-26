"""Small helpers over python-docx for building the review document.

The course specifies Times New Roman at 12 pt, so that is enforced on the base
style rather than applied run by run, which is easy to get inconsistently wrong.
"""

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(12)
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x3B, 0x73)
MUTED = RGBColor(0x55, 0x55, 0x66)


def new_document() -> Document:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    # East-Asian font name must be set separately or Word may substitute a font.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for name, size, before, after in [
        ("Heading 1", 15.5, 10, 6),
        ("Heading 2", 13, 9, 4),
        ("Heading 3", 12, 7, 3),
    ]:
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    return doc


def add_page_numbers(doc: Document, label: str) -> None:
    """Footer with 'label | Page N of M'. Needs raw field codes."""
    footer = doc.sections[-1].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run(f"{label}    |    Page ")
    run.font.size = Pt(9)
    run.font.name = BODY_FONT

    def field(instr: str) -> None:
        r = p.add_run()
        r.font.size = Pt(9)
        r.font.name = BODY_FONT
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = instr
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        r._r.append(begin)
        r._r.append(instr_el)
        r._r.append(end)

    field("PAGE")
    mid = p.add_run(" of ")
    mid.font.size = Pt(9)
    mid.font.name = BODY_FONT
    field("NUMPAGES")


def shade(cell_or_para, hex_colour: str) -> None:
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_colour)
    target = cell_or_para._tc if hasattr(cell_or_para, "_tc") else cell_or_para._p.get_or_add_pPr()
    target.append(el)


def h1(doc: Document, text: str, page_break: bool = False) -> None:
    """Sections run on by default. At 12 pt the page budget cannot absorb eleven
    part-title pages, so only the front matter forces a break."""
    if page_break and len(doc.paragraphs) > 1:
        doc.add_page_break()
    doc.add_heading(text, level=1)


def h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def h3(doc: Document, text: str) -> None:
    doc.add_heading(text, level=3)


def para(doc: Document, text: str, justify: bool = True, italic: bool = False,
         size: float | None = None) -> None:
    p = doc.add_paragraph()
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.italic = italic
    if size:
        run.font.size = Pt(size)


def bullets(doc: Document, items: list[str], style: str = "List Bullet") -> None:
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        # "**bold**:" prefix marks the term being defined.
        if item.startswith("**") and "**" in item[2:]:
            term, rest = item[2:].split("**", 1)
            r = p.add_run(term)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def code(doc: Document, text: str, size: float = 7.6, shade_hex: str = "F3F5F9") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.0
    shade(p, shade_hex)
    run = p.add_run(text.rstrip())
    run.font.name = CODE_FONT
    run.font.size = Pt(size)
    run.element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)


def compact_json(value, width: int = 96, indent: int = 2, _level: int = 0) -> str:
    """Pretty-print JSON, but keep any subtree that fits on one line on one line.

    json.dumps(indent=2) puts every array element on its own row, which turns a
    short boundary list or a $cond into a dozen nearly empty lines. Collapsing
    subtrees that fit lets a whole pipeline be shown instead of a truncated one.
    """
    pad = " " * (indent * _level)
    flat = json.dumps(value, default=str)
    if len(flat) + len(pad) <= width or not isinstance(value, (dict, list)):
        return pad + flat

    inner = " " * (indent * (_level + 1))
    if isinstance(value, dict):
        parts = []
        for k, v in value.items():
            rendered = compact_json(v, width, indent, _level + 1).lstrip()
            parts.append(f"{inner}{json.dumps(k)}: {rendered}")
        return pad + "{\n" + ",\n".join(parts) + "\n" + pad + "}"

    parts = [compact_json(v, width, indent, _level + 1) for v in value]
    return pad + "[\n" + ",\n".join(parts) + "\n" + pad + "]"


def json_block(doc: Document, value, size: float = 7.2, max_chars: int = 1500) -> None:
    text = compact_json(value)
    if len(text) > max_chars:
        text = text[:max_chars].rsplit("\n", 1)[0] + "\n  ... truncated ..."
    code(doc, text, size=size, shade_hex="FBF7EF")


def table(doc: Document, headers: list[str], rows: list[list], widths: list[float] | None = None,
          font_size: float = 9.0, style: str = "Table Grid") -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.name = BODY_FONT
        shade(cell, "DCE4F2")

    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run("" if value is None else str(value))
            run.font.size = Pt(font_size)
            run.font.name = BODY_FONT

    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)

    # A spacer keeps consecutive tables from merging, but a full 12 pt empty line
    # after each of two dozen tables costs most of a page, so keep it tiny.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.add_run("").font.size = Pt(4)
    return t


def figure(doc: Document, path: Path, caption: str, width: float = 6.2) -> None:
    if not path.exists():
        para(doc, f"[missing figure: {path.name}]", italic=True)
        return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(caption)
    run.italic = True
    run.font.size = Pt(9.5)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)


def figure_grid(doc: Document, items: list[tuple[Path, str]], width: float = 3.0,
                columns: int = 2) -> None:
    """Lay screenshots out several to a row inside a borderless table.

    One figure per row wastes the right half of the page and costs over a page
    across a gallery of six, which the page budget cannot afford.
    """
    items = [(p, c) for p, c in items if p.exists()]
    if not items:
        return

    rows = (len(items) + columns - 1) // columns
    t = doc.add_table(rows=rows, cols=columns)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    for n, (path, caption) in enumerate(items):
        cell = t.rows[n // columns].cells[n % columns]
        cell.width = Inches(width + 0.15)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        p.add_run().add_picture(str(path), width=Inches(width))

        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.space_after = Pt(6)
        run = cp.add_run(caption)
        run.italic = True
        run.font.size = Pt(8.5)
        run.font.name = BODY_FONT

    # Trailing cells in a short final row must stay empty and unobtrusive.
    for n in range(len(items), rows * columns):
        t.rows[n // columns].cells[n % columns].paragraphs[0].add_run("").font.size = Pt(4)


def caption_only(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(10)


def landscape_section(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = 1
    section.page_width, section.page_height = section.page_height, section.page_width
